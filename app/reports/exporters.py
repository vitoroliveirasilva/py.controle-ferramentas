from __future__ import annotations

import csv
from io import BytesIO, StringIO
from typing import Any

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

FORMULA_PREFIXES = ("=", "+", "-", "@")


def sanitize_export_value(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, str):
        cleaned = value.replace("\x00", "")
        if cleaned.startswith(FORMULA_PREFIXES):
            return f"'{cleaned}"
        return cleaned
    return value


def sanitize_export_rows(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    return [
        {key: sanitize_export_value(value) for key, value in row.items()}
        for row in rows
    ]


def rows_to_csv(rows: list[dict[str, Any]]) -> str:
    output = StringIO()
    if not rows:
        return ""
    sanitized_rows = sanitize_export_rows(rows)
    writer = csv.DictWriter(
        output,
        fieldnames=list(sanitized_rows[0].keys()),
        extrasaction="ignore",
    )
    writer.writeheader()
    writer.writerows(sanitized_rows)
    return output.getvalue()


def rows_to_xlsx(rows: list[dict[str, Any]], sheet_name: str = "Relatorio") -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet_name[:31] or "Relatorio"

    if rows:
        sanitized_rows = sanitize_export_rows(rows)
        headers = list(sanitized_rows[0].keys())
        worksheet.append(headers)
        for row in sanitized_rows:
            worksheet.append([row.get(header) for header in headers])
        for column_index, header in enumerate(headers, start=1):
            max_length = max(
                len(str(header)),
                *(len(str(row.get(header, "") or "")) for row in sanitized_rows),
            )
            worksheet.column_dimensions[get_column_letter(column_index)].width = min(
                max(12, max_length + 2),
                48,
            )
    else:
        worksheet.append(["mensagem"])
        worksheet.append(["Nenhum registro encontrado para os filtros informados."])

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
